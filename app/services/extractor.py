from __future__ import annotations

import re
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

MAX_PDF_PAGES = 300


class ExtractionError(Exception):
    """A document could not be read in a way the user can act on."""


@dataclass
class ExtractedDocument:
    source_type: str
    markdown: str
    as_json: dict[str, Any]
    blocks: list[dict[str, Any]]


def extract_document(path: Path) -> ExtractedDocument:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    raise ExtractionError("仅支持 .docx 和 .pdf 文件。")


def extract_docx(path: Path) -> ExtractedDocument:
    try:
        document = Document(str(path))
    except Exception as exc:
        raise ExtractionError("Word 文件无法打开，请确认它是有效的 .docx 文件。") from exc

    blocks: list[dict[str, Any]] = []
    markdown_parts: list[str] = []

    try:
        for item in _iter_block_items(document):
            if isinstance(item, Paragraph):
                text = item.text.replace("\u00a0", " ").strip()
                if not text:
                    continue
                style = item.style.name if item.style is not None else "Normal"
                level = _heading_level(style)
                blocks.append(
                    {
                        "type": "paragraph",
                        "text": text,
                        "style": style,
                        **({"heading_level": level} if level else {}),
                    }
                )
                if level:
                    markdown_parts.append(f"{'#' * min(level + 1, 6)} {text}")
                elif _is_bullet(style):
                    markdown_parts.append(f"- {text}")
                else:
                    markdown_parts.append(text)
            elif isinstance(item, Table):
                rows = _table_rows(item)
                if not rows:
                    continue
                blocks.append({"type": "table", "rows": rows})
                markdown_parts.append(_rows_to_markdown(rows))
    except Exception as exc:
        raise ExtractionError("读取 Word 内容时出错，文件可能已损坏。") from exc

    if not blocks:
        raise ExtractionError("文档中没有可读取的文字内容。")

    markdown = "\n\n".join(markdown_parts).strip() + "\n"
    table_count = sum(1 for block in blocks if block["type"] == "table")
    json_data = {
        "source_type": "docx",
        "blocks": blocks,
        "stats": {"blocks": len(blocks), "tables": table_count},
    }
    return ExtractedDocument("docx", markdown, json_data, blocks)


def extract_pdf(path: Path) -> ExtractedDocument:
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractionError("PDF 已加密，请先取消密码保护后再上传。") from exc
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ExtractionError(f"PDF 页数超过 {MAX_PDF_PAGES} 页，请拆分后再上传。")

        blocks: list[dict[str, Any]] = []
        pages: list[dict[str, Any]] = []
        markdown_parts: list[str] = []
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                raw_text = page.extract_text() or ""
            except Exception:  # noqa: BLE001 - one unreadable page should not fail the PDF
                raw_text = ""
            text = raw_text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ").strip()
            paragraphs = _split_pdf_text(text)
            pages.append({"page": page_number, "text": text})
            if not paragraphs:
                continue
            for paragraph in paragraphs:
                blocks.append({"type": "paragraph", "text": paragraph, "page": page_number})
            markdown_parts.append(f"## 第 {page_number} 页")
            markdown_parts.extend(paragraphs)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("PDF 文件无法打开，请确认文件有效且未损坏。") from exc

    if not pages or not any(page["text"] for page in pages):
        raise ExtractionError("PDF 中没有可读取的文字，可能是扫描件。请先做 OCR 后再上传。")

    markdown = "\n\n".join(markdown_parts).strip() + "\n"
    json_data = {
        "source_type": "pdf",
        "pages": pages,
        "blocks": blocks,
        "stats": {"pages": len(pages), "blocks": len(blocks)},
    }
    return ExtractedDocument("pdf", markdown, json_data, blocks)


def _iter_block_items(parent: DocxDocument) -> Iterator[Paragraph | Table]:
    parent_element = parent.element.body
    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _table_rows(table: Table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            text_parts = []
            for paragraph in cell.paragraphs:
                text = paragraph.text.replace("\u00a0", " ").strip()
                if text:
                    text_parts.append(text)
            cells.append("\n".join(text_parts))
        rows.append(cells)
    return rows


def _rows_to_markdown(rows: list[list[str]]) -> str:
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    header = padded[0]
    divider = ["---"] * width
    body = padded[1:]
    lines = [
        "| " + " | ".join(_escape_cell(cell) for cell in header) + " |",
        "| " + " | ".join(divider) + " |",
    ]
    lines.extend("| " + " | ".join(_escape_cell(cell) for cell in row) + " |" for row in body)
    return "\n".join(lines)


def _escape_cell(value: str) -> str:
    return value.replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def _split_pdf_text(text: str) -> list[str]:
    if not text:
        return []
    # Blank lines are reliable paragraph boundaries. If a PDF has no blank
    # lines, each extracted line is kept as its own block so equations survive.
    chunks = re.split(r"\n\s*\n", text)
    paragraphs: list[str] = []
    for chunk in chunks:
        lines = [line.strip() for line in chunk.splitlines() if line.strip()]
        if not lines:
            continue
        if len(chunks) > 1:
            paragraphs.append(" ".join(lines))
        else:
            paragraphs.extend(lines)
    return paragraphs


def _heading_level(style_name: str) -> int | None:
    match = re.match(r"^(?:Heading|标题)\s*(\d+)$", style_name.strip(), re.IGNORECASE)
    if not match:
        return None
    return max(1, min(int(match.group(1)), 6))


def _is_bullet(style_name: str) -> bool:
    return style_name.lower().startswith(("list bullet", "列表项目符号"))