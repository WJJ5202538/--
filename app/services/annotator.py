from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.extractor import ExtractedDocument
from app.services.mathcheck import MathIssue, find_math_issues

ANNOTATION_COLOR = RGBColor(0xC0, 0x39, 0x2B)
ANNOTATION_FONT = "Microsoft YaHei"


class AnnotationError(Exception):
    """The annotated document could not be written."""


def annotate_docx(source: Path, output: Path) -> int:
    """Add red # annotations to a copy of the original docx.

    Existing text is never edited. Each annotation is inserted as a new
    paragraph directly after the paragraph that contains the mistake.
    """
    try:
        document = Document(str(source))
    except Exception as exc:
        raise AnnotationError("Word 文件无法读取，不能生成批注稿。") from exc

    issue_count = 0
    try:
        for paragraph in _iter_paragraphs(document):
            issues = find_math_issues(paragraph.text)
            if issues:
                _insert_annotations(paragraph, issues)
                issue_count += len(issues)
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output))
    except AnnotationError:
        raise
    except Exception as exc:
        raise AnnotationError("写入批注稿时出错，请检查文件是否被 Word 占用。") from exc

    return issue_count


def create_annotated_docx_from_pdf(extracted: ExtractedDocument, output: Path) -> int:
    """Build an editable review copy for a PDF upload.

    PDFs cannot be edited in place like DOCX files. The extracted text is
    reflowed into a clean Word document, then annotated in the same way.
    """
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = document.styles["Normal"]
    normal.font.name = ANNOTATION_FONT
    normal.font.size = Pt(10.5)
    _set_east_asia_font(normal.element, ANNOTATION_FONT)

    blocks = extracted.as_json.get("blocks", [])
    pages = extracted.as_json.get("pages", [])
    issue_count = 0
    last_page = None

    try:
        for page in pages:
            page_number = page.get("page")
            page_blocks = [block for block in blocks if block.get("page") == page_number]
            if not page_blocks:
                continue
            if last_page is not None:
                document.add_page_break()
            last_page = page_number

            for block in page_blocks:
                text = str(block.get("text", "")).strip()
                if not text:
                    continue
                paragraph = document.add_paragraph(text)
                paragraph.paragraph_format.space_after = Pt(6)
                paragraph.paragraph_format.line_spacing = 1.5
                issues = find_math_issues(text)
                if issues:
                    _insert_annotations(paragraph, issues)
                    issue_count += len(issues)

        if last_page is None:
            raise AnnotationError("PDF 中没有可写入批注稿的文字。")

        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output))
    except AnnotationError:
        raise
    except Exception as exc:
        raise AnnotationError("生成 PDF 批注稿时出错。") from exc

    return issue_count


def _iter_paragraphs(document) -> Iterator[Paragraph]:
    # Track the lxml elements themselves. Using id(element) here is unsafe:
    # temporary python-docx wrappers can be collected and their ids reused,
    # which silently skips table cells.
    seen: set[object] = set()

    def emit(paragraphs):
        for paragraph in paragraphs:
            element = paragraph._p
            if element not in seen:
                seen.add(element)
                yield paragraph

    yield from emit(document.paragraphs)
    for table in document.tables:
        yield from _iter_table_paragraphs(table, seen)


def _iter_table_paragraphs(table: Table, seen: set[object]) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                element = paragraph._p
                if element not in seen:
                    seen.add(element)
                    yield paragraph
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested, seen)


def _insert_annotations(paragraph: Paragraph, issues: list[MathIssue]) -> None:
    anchor = paragraph._p
    for issue in issues:
        new_element = OxmlElement("w:p")
        anchor.addnext(new_element)
        annotation = Paragraph(new_element, paragraph._parent)
        annotation.alignment = WD_ALIGN_PARAGRAPH.LEFT
        annotation.paragraph_format.space_before = Pt(0)
        annotation.paragraph_format.space_after = Pt(4)
        annotation.paragraph_format.left_indent = Pt(0)
        run = annotation.add_run(issue.message)
        _style_annotation_run(run)
        anchor = new_element


def _style_annotation_run(run) -> None:
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = ANNOTATION_COLOR
    run.font.name = ANNOTATION_FONT
    _set_east_asia_font(run._element, ANNOTATION_FONT)


def _set_east_asia_font(element, font_name: str) -> None:
    r_pr = element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)