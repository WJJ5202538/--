from pathlib import Path

from docx import Document
from helpers import docx_bytes, pdf_bytes

from app.services.annotator import annotate_docx, create_annotated_docx_from_pdf
from app.services.extractor import extract_docx, extract_pdf


def test_extracts_paragraphs_and_tables_to_json_and_markdown(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    source.write_bytes(docx_bytes(["第一段", "第二段"], [["项目", "结果"], ["3 + 4 = 8", "错误"]]))

    extracted = extract_docx(source)

    assert extracted.source_type == "docx"
    assert extracted.as_json["stats"]["tables"] == 1
    assert "第一段" in extracted.markdown
    assert "| 项目 | 结果 |" in extracted.markdown
    assert "3 + 4 = 8" in extracted.markdown


def test_annotates_docx_without_changing_original_text(tmp_path: Path) -> None:
    source = tmp_path / "original.docx"
    source.write_bytes(
        docx_bytes(
            ["计算：3 × 4 = 13。", "正确：2 + 2 = 4。"],
            [["表格题", "10 ÷ 2 = 6"], ["答案", "5"]],
        )
    )
    output = tmp_path / "annotated.docx"

    count = annotate_docx(source, output)
    assert count == 2
    assert output.is_file()

    document = Document(output)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert "计算：3 × 4 = 13。" in paragraphs
    assert "正确：2 + 2 = 4。" in paragraphs
    assert sum(text.startswith("#") for text in paragraphs) == 1
    index = paragraphs.index("计算：3 × 4 = 13。")
    assert paragraphs[index + 1] == "# 计算有误：3 × 4 = 13，正确结果应为 12。"
    annotation_run = document.paragraphs[index + 1].runs[0]
    assert annotation_run.bold is True
    assert str(annotation_run.font.color.rgb) == "C0392B"

    checked_cell = document.tables[0].rows[0].cells[1]
    assert checked_cell.paragraphs[0].text == "10 ÷ 2 = 6"
    assert checked_cell.paragraphs[1].text == "# 计算有误：10 ÷ 2 = 6，正确结果应为 5。"

    source_document = Document(source)
    assert [paragraph.text for paragraph in source_document.paragraphs] == [
        "计算：3 × 4 = 13。",
        "正确：2 + 2 = 4。",
    ]


def test_pdf_extraction_and_annotated_word_generation(tmp_path: Path) -> None:
    source = tmp_path / "sample.pdf"
    source.write_bytes(pdf_bytes("3 + 5 = 9"))
    extracted = extract_pdf(source)
    assert extracted.source_type == "pdf"
    assert "3 + 5 = 9" in extracted.markdown
    assert extracted.as_json["stats"]["pages"] == 1

    output = tmp_path / "review.docx"
    count = create_annotated_docx_from_pdf(extracted, output)
    assert count == 1
    texts = [paragraph.text for paragraph in Document(output).paragraphs]
    assert "3 + 5 = 9" in texts
    assert "# 计算有误：3 + 5 = 9，正确结果应为 8。" in texts